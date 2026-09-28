"""
Ingest 3 Gemini C7 womenwork ablation cells (sen, cot, pc) into calibration_final/:
  - Raw JSONL → stages/stage_4_prompt_design/raw_jsonl/womenwork/
  - Metrics    → stages/stage_4_prompt_design/metrics/{ablation}_gemini_vs_T08_metrics.csv
  - Excel row  → per_outcome/womenwork/womenwork_all_phases.xlsx (Stage 4, Phase 6+13/14/15)
Idempotent: skip cells already present.
"""
from __future__ import annotations
import json, shutil
from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd
from scipy.spatial.distance import jensenshannon
from scipy.stats import wasserstein_distance
from sklearn.metrics import cohen_kappa_score

ROOT = Path("/Users/helinekmen/Desktop/digitaltwin")
CACHE = Path.home() / "Library/Caches/digitaltwin_calibration"
FINAL = ROOT / "outputs/calibration_final"
GT_CSV = ROOT / "agent-survey/agent-calibration/womenwork/womenwork-data/womenwork_data_output.csv"

# (cache filename, ablation label, phase label, metrics CSV filename)
CELLS = [
    ("C7_geminiflashlite_T08_sen_vs_cot_nocot_s0.jsonl",
     "sen-dili address mode (final pick)",
     "Phase 6+13", "sendili_geminiflashlite_vs_T08_metrics.csv"),
    ("C7_geminiflashlite_T08_ben_vs_cot_cot_s0.jsonl",
     "Explicit CoT (final pick)",
     "Phase 14",   "explicitcot_geminiflashlite_vs_T08_metrics.csv"),
    ("C7_geminiflashlite_T08_ben_vs_cot_nocot_pc_s0.jsonl",
     "Political context prepend (final pick)",
     "Phase 15",   "politicalcontext_geminiflashlite_vs_T08_metrics.csv"),
]

STAGE4_RAW = FINAL / "stages/stage_4_prompt_design/raw_jsonl/womenwork"
STAGE4_MET = FINAL / "stages/stage_4_prompt_design/metrics"
STAGE4_RAW.mkdir(parents=True, exist_ok=True)

# Load GT
g = pd.read_csv(GT_CSV)
g["rid"] = g["persona_id"].astype(str).str.zfill(4).apply(lambda x: f"TGSS_{x}")
gts = dict(zip(g["rid"], g["gt_womenwork"].astype(int)))


def load_preds(fp):
    out = {}
    for line in fp.open(encoding="utf-8"):
        try: r = json.loads(line)
        except: continue
        if r.get("parse_status") == "ok" and r.get("predicted_value") is not None:
            out[r["respondent_id"]] = int(r["predicted_value"])
    return out


def dist(vals, lo, hi):
    counts = np.zeros(hi - lo + 1)
    for v in vals:
        if lo <= v <= hi: counts[v - lo] += 1
    return counts / max(counts.sum(), 1)


def compute(fp):
    preds = load_preds(fp)
    pairs = [(gts[r], preds[r]) for r in preds if r in gts]
    if not pairs: return None
    gt = np.array([p[0] for p in pairs]); pr = np.array([p[1] for p in pairs])
    labels = [1,2,3,4,5]
    gt_p, pr_p = dist(gt,1,5), dist(pr,1,5)
    return dict(
        n_matched=len(pairs),
        wasserstein=float(wasserstein_distance([1,2,3,4,5],[1,2,3,4,5],gt_p,pr_p)),
        jsd=float(jensenshannon(gt_p, pr_p, base=2)**2),
        tvd=float(np.abs(gt_p-pr_p).sum()/2),
        pred_mean=float(pr.mean()), gt_mean=float(gt.mean()),
        prevalence_error=float(abs(gt.mean()-pr.mean())),
        cohen_kappa=float(cohen_kappa_score(gt, pr, labels=labels)),
        weighted_kappa_quadratic=float(cohen_kappa_score(gt, pr, labels=labels, weights="quadratic")),
        weighted_kappa_linear=float(cohen_kappa_score(gt, pr, labels=labels, weights="linear")),
    )


excel_fp = FINAL / "per_outcome/womenwork/womenwork_all_phases.xlsx"
wb = openpyxl.load_workbook(excel_fp)
ws = wb["All cells"]
header = [c.value for c in ws[1]]

# Existing keys for dup check
existing = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    rec = dict(zip(header, row))
    existing.add((rec.get("Config"), rec.get("Model"), rec.get("Sampling"),
                  rec.get("Temperature"), rec.get("Address mode"),
                  rec.get("CoT"), rec.get("Seed"), rec.get("Phase"),
                  # Include a marker for political_context in Notes
                  "pc" if rec.get("Notes") and "political context" in str(rec.get("Notes","")).lower() else ""))

copied = metrics_added = excel_added = 0
for cache_name, ablation_label, phase_label, csv_name in CELLS:
    src = CACHE / "screening_vs_cot_T08/womenwork" / cache_name
    if not src.exists():
        print(f"  ⚠ MISSING: {cache_name}"); continue

    # 1) raw copy
    dst = STAGE4_RAW / cache_name
    if not dst.exists() or dst.stat().st_size != src.stat().st_size:
        shutil.copy2(src, dst); copied += 1

    # 2) metrics
    m = compute(src)
    if m is None:
        print(f"  ⚠ no pairs: {cache_name}"); continue

    # Determine cell properties from filename
    address = "sen" if "_sen_" in cache_name else "ben"
    cot = "_cot_s0" in cache_name
    pc = "_pc_s0" in cache_name

    # 3) metrics CSV
    csv_fp = STAGE4_MET / csv_name
    row = dict(
        outcome="womenwork", config_id="C7", step="2A_screening",
        model="gemini-2.5-flash-lite", temperature=0.8, address_mode=address,
        sampling="vs_cot", cot=cot, seed=0,
        n_total=m["n_matched"], n_matched=m["n_matched"],
        n_parse_failures=0, n_api_errors=0,
        jsd=m["jsd"], wasserstein=m["wasserstein"], tvd=m["tvd"],
        pred_mean=m["pred_mean"], gt_mean=m["gt_mean"],
        prevalence_diff=m["pred_mean"] - m["gt_mean"],
        prevalence_error=m["prevalence_error"],
        precision_minority=None, recall_minority=None, f1_minority=None,
        macro_f1=None, mcc=None, cohen_kappa=m["cohen_kappa"],
        composite_score=None,
        timestamp=pd.Timestamp.now(tz="UTC").isoformat(),
        political_context=pc,
    )
    if csv_fp.exists():
        df = pd.read_csv(csv_fp)
        mask = ((df["outcome"]=="womenwork") & (df["config_id"]=="C7")
                & (df["address_mode"]==address) & (df["cot"]==cot))
        if "political_context" in df.columns:
            mask &= (df["political_context"] == pc)
        if mask.any():
            print(f"    {csv_name}: already present, skipping")
        else:
            df = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
            df.to_csv(csv_fp, index=False); metrics_added += 1
            print(f"    ✓ {csv_name}: appended")
    else:
        pd.DataFrame([row]).to_csv(csv_fp, index=False); metrics_added += 1
        print(f"    ✓ {csv_name}: created")

    # 4) Excel row
    excel_key = ("C7", "gemini-2.5-flash-lite", "vs", 0.8, address,
                 cot, 0, phase_label, "pc" if pc else "")
    if excel_key in existing:
        print(f"    Excel: {excel_key} already present, skipping")
        continue
    excel_row = {
        "Phase": phase_label,
        "Stage": "Stage 4 — Prompt design",
        "Section": {"Phase 6+13": "S06+S13", "Phase 14": "S14",
                    "Phase 15": "S15"}.get(phase_label, ""),
        "Experiment": f"{ablation_label} on final pick (C7·Gemini FL VS T=0.8)",
        "Config": "C7", "Model": "gemini-2.5-flash-lite",
        "Sampling": "vs", "Temperature": 0.8,
        "Address mode": address, "CoT": cot, "Seed": 0,
        "n_matched": m["n_matched"],
        "wasserstein": m["wasserstein"], "jsd": m["jsd"],
        "tvd": m["tvd"], "pred_mean": m["pred_mean"], "gt_mean": m["gt_mean"],
        "prevalence_error": m["prevalence_error"],
        "Notes": (f"Stage 4 ablation re-run on the FINAL PICK cell "
                  f"(C7·Gemini FL·VS·T=0.8) — supersedes the earlier gpt-4o-mini "
                  f"ablations for this stage's decisions."
                  + (" [political context prepend]" if pc else "")
                  + (" [explicit CoT]" if cot else "")
                  + (" [sen-dili address mode]" if address == "sen" else "")),
        "Source_CSV": f"stages/stage_4_prompt_design/metrics/{csv_name}",
        "cohen_kappa": m["cohen_kappa"],
        "weighted_kappa_quadratic": m["weighted_kappa_quadratic"],
        "weighted_kappa_linear": m["weighted_kappa_linear"],
        "ac2_quadratic": None,
    }
    ws.append([excel_row.get(h) for h in header])
    excel_added += 1

wb.save(excel_fp)

print()
print(f"✓ SUMMARY")
print(f"  Raw JSONL copied to stage_4:   {copied}")
print(f"  Metrics CSV rows appended:     {metrics_added}")
print(f"  Excel rows added:              {excel_added}  →  {excel_fp.relative_to(ROOT)}")

# Quick comparison table
print("\n=== C7 · Gemini FL · VS T=0.8 — ablation comparison ===")
baseline_wass, baseline_kq = 0.218, 0.401
print(f"  {'Ablation':30s}  {'Wass':>7s}  {'Δ':>7s}  {'κ_quad':>7s}  {'Δ':>7s}")
print(f"  {'BASELINE (no ablation)':30s}  {baseline_wass:>7.3f}  {'—':>7s}  {baseline_kq:>7.3f}  {'—':>7s}")
for cache_name, ablation_label, _, _ in CELLS:
    src = CACHE / "screening_vs_cot_T08/womenwork" / cache_name
    m = compute(src)
    if m is None: continue
    dW = m['wasserstein'] - baseline_wass
    dK = m['weighted_kappa_quadratic'] - baseline_kq
    print(f"  {ablation_label:30s}  {m['wasserstein']:>7.3f}  {dW:>+7.3f}  "
          f"{m['weighted_kappa_quadratic']:>7.3f}  {dK:>+7.3f}")
